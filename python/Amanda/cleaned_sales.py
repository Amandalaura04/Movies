# Cleaning the movie datasets
#
# I use this script to clean the sales file and prepare the IMDb movie data.
# The original files and cleaned outputs are all in my Downloads folder.
#
# Amanda Luijendijk | [private student number]
# Database Management (7000DMD_23), 2026-2027

from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

DOWNLOADS = Path.home() / "Downloads"

# Some movie names have strange charachters because of how the text was imported.
MOJIBAKE = re.compile(r"[ÃÂ]|â€")

MONTHS = ("January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December")


# 1. Small cleaning functions

def fix_mojibake(s):
    if not isinstance(s, str) or not s:
        return s
    for _ in range(3):
        if not MOJIBAKE.search(s):
            break
        try:
            fixed = s.encode("cp1252").decode("utf-8")
        except (UnicodeDecodeError, UnicodeEncodeError):
            break
        if fixed == s:
            break
        s = fixed
    return s


def clean_title(row):
    t = row["title"]
    if isinstance(t, str):
        return t.strip()
    if isinstance(t, (int, float, np.integer, np.floating)) and not pd.isna(t):
        return str(int(t)) if float(t).is_integer() else str(t)

    # Excel sometimes reads movie titles as dates. I use the URL to get the title back.
    url = row["url"]
    if not isinstance(url, str) or "/movie/" not in url:
        return pd.NA
    slug = url.rsplit("/movie/", 1)[-1]
    slug = re.sub(r"\(.*?\)", "", slug)
    return re.sub(r"\s+", " ", slug.replace("-", " ")).strip()


def parse_release_date(row):
    rd, year = row["release_date"], row["year"]
    if not isinstance(rd, str) or not rd.startswith(MONTHS) or pd.isna(year):
        return pd.NaT
    rd = re.sub(r"(\d+)(st|nd|rd|th)", r"\1", rd).strip()
    return pd.to_datetime(f"{rd} {int(year)}", format="%B %d %Y", errors="coerce")


# 2. Clean the sales dataset

def clean_sales(path):
    df = pd.read_excel(path, sheet_name="numbers253")
    n_start = len(df)

    df = df.drop(columns=["Unnamed: 8"])
    df["title"] = df.apply(clean_title, axis=1)
    df["release_date"] = df.apply(parse_release_date, axis=1)

    for col in ["title", "genre", "keywords", "creative_type"]:
        df[col] = df[col].apply(fix_mojibake).astype("string").str.strip()
    df["url"] = df["url"].astype("string").str.strip()

    df = df[df["title"].notna() & (df["title"].str.len() > 0)]

    df = df[["year", "release_date", "title", "genre",
             "international_box_office", "domestic_box_office",
             "worldwide_box_office", "production_budget", "opening_weekend",
             "theatre_count", "avg run per theatre", "runtime", "keywords",
             "creative_type", "url"]]

    print(f"Sales: {n_start} -> {len(df)} rows, "
          f"{df['release_date'].isna().sum()} without a release date")
    return df


# 3. Prepare the IMDb dataset

def find_file(*names):
    for name in names:
        hits = list(DOWNLOADS.glob(name))
        if hits:
            return hits[0]
    raise FileNotFoundError(f"None of {names} found in {DOWNLOADS}")


def clean_imdb(basics_path, ratings_path):
    gz = lambda p: "gzip" if str(p).endswith(".gz") else None

    # This file is quite large, so I read it in smaller chuncks and keep the movies.
    chunks = []
    for chunk in pd.read_csv(basics_path, sep="\t", compression=gz(basics_path),
                             na_values="\\N", dtype=str, chunksize=500_000):
        chunks.append(chunk[chunk["titleType"].isin(["movie", "tvMovie"])])
    basics = pd.concat(chunks, ignore_index=True)

    ratings = pd.read_csv(ratings_path, sep="\t", compression=gz(ratings_path),
                          na_values="\\N")

    df = basics.merge(ratings, on="tconst", how="left")
    df = df.drop(columns=["endYear"])
    # astype(bool) turns the text "0" into True as well, so compare with "1"
    df["isAdult"] = df["isAdult"].eq("1")
    df["startYear"] = df["startYear"].astype("Int64")
    df["runtimeMinutes"] = df["runtimeMinutes"].astype("Int64")
    df = df.rename(columns={
        "primaryTitle": "title",
        "originalTitle": "original_title",
        "startYear": "year",
        "runtimeMinutes": "runtime",
        "averageRating": "imdb_rating",
        "numVotes": "imdb_votes",
    })

    # Put IMDb ratings on a 0-100 scale for a consistant comparison.
    df["imdb_rating"] = (df["imdb_rating"] * 10).round(1)

    print(f"IMDb: {len(df)} movies, {df['imdb_rating'].notna().sum()} with a rating")
    return df


# 4. Save the cleaned data

def save_excel(df, path, sheet):
    df.to_excel(path, sheet_name=sheet, index=False)
    wb = load_workbook(path)
    ws = wb[sheet]
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=cell.row == 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# 5. Run the cleaning steps

# Save both cleaned datasets in Downloads, so I can find them laterr.
if __name__ == "__main__":
    sales = clean_sales(DOWNLOADS / "sales.xlsx")
    save_excel(sales, DOWNLOADS / "sales_cleaned.xlsx", "sales")

    basics = find_file("title.basics.tsv", "title_basics_tsv.gz", "title.basics.tsv.gz")
    ratings = find_file("title.ratings.tsv", "title_ratings_tsv.gz", "title.ratings.tsv.gz")
    imdb = clean_imdb(basics, ratings)
    imdb.to_csv(DOWNLOADS / "imdb_movies_cleaned.csv", index=False)

    print(f"Saved sales_cleaned.xlsx and imdb_movies_cleaned.csv in {DOWNLOADS}")
